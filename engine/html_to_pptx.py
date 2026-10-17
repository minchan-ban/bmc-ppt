#!/usr/bin/env python3
"""HTML 덱(16:9, 1 section = 1 슬라이드)을 편집 가능한 PPTX로 변환한다.

사용법:
  .venv/bin/python engine/html_to_pptx.py deck.html              # 같은 폴더에 deck.pptx
  .venv/bin/python engine/html_to_pptx.py deck.html out.pptx

원본: carnot-tech/consulting-pptx-skill scripts/html_to_pptx.py (MIT). bmc-ppt용 수정 사항:
  - 글꼴을 Pretendard로 고정하고, 굵기를 굵기별 글꼴 이름으로 지정한다 (WEIGHT_FACE 참조)
  - 헤드메시지를 제목 개체 틀에 넣는다 (PowerPoint 개요 보기와 슬라이드 탐색에 제목이 나온다)
  - 목록(ul/ol)을 문단이 여러 개인 글상자 하나로 만든다 (글머리표를 이어서 편집할 수 있다)
  - 블록 글상자는 원래 블록 너비로 만들어 자동 줄바꿈을 켠다 (글자를 늘려도 배치가 유지된다)
  - 테마 글꼴과 테마 색을 덱의 토큰으로 바꾼다 (새로 그리는 도형과 글상자도 같은 체계를 따른다)
  - 한글을 어절 단위로 줄바꿈한다 (낱말 중간에서 줄이 끊기지 않는다)

변환 방식: html_dump.mjs가 Chrome으로 그린 요소의 위치와 모양을 기록하고, 이 스크립트가 PowerPoint 도형으로 조립한다.
  글자 → 글상자 / 채움·테두리 → 사각형·다각형·타원 / 한쪽 괘선 → 직선 / 표 → PowerPoint 표
  SVG·img·배경 이미지 → PNG 그림 (안의 글자와 수치는 편집할 수 없으므로 부품에서 SVG를 쓰지 않는다)
  .nchart (원형·도넛·꺾은선·영역) → PowerPoint 기본 차트 (engine/native_chart.py. '데이터 편집'으로 수치를 고친다)
필요한 것: Node 22 이상, Chrome, python-pptx.
"""
import json
import re
import subprocess
import sys
import tempfile
from pathlib import Path
from urllib.parse import unquote, urlparse

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
from native_chart import add_chart  # noqa: E402  원형·도넛·꺾은선·영역을 PowerPoint 기본 차트로

SLIDE_W, SLIDE_H = 12192000, 6858000
FAMILY = "Pretendard"
# CSS font-weight → (PowerPoint 글꼴 이름, 굵게 여부).
# Regular와 Bold만 같은 글꼴 이름을 쓰고, 나머지 굵기는 굵기 이름이 붙은 별도 글꼴 이름을 쓴다.
# 이렇게 지정해야 macOS와 Windows의 PowerPoint 모두에서 굵기가 그대로 보이고, 글꼴 목록에서 바꿀 수 있다.
WEIGHT_FACE = {
    300: ("Pretendard Light", False),
    400: ("Pretendard", False),
    500: ("Pretendard Medium", False),
    600: ("Pretendard SemiBold", False),
    700: ("Pretendard", True),
    800: ("Pretendard ExtraBold", False),
    900: ("Pretendard Black", False),
}
NO_STYLE_TABLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # 표 스타일 "스타일 없음, 눈금 없음"
ALIGN = {"left": PP_ALIGN.LEFT, "start": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER,
         "right": PP_ALIGN.RIGHT, "end": PP_ALIGN.RIGHT, "justify": PP_ALIGN.LEFT}
VALIGN = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}
SHAPE_NAMES = {"title": "헤드메시지", "colh": "열 제목", "axh": "축 제목", "bl": "글머리표", "src": "출처", "foot": "꼬리말",
               "kicker": "키커", "brand": "회사명", "bar": "막대", "cb": "막대", "cv": "단계", "sg": "구성 조각", "kpi": "핵심 수치",
               "t": "표", "sp": "주장 패널", "tri": "인과 화살표", "hr": "제목 괘선", "pt": "위치 점", "hb": "충족도 기호"}


class Scale:
    def __init__(self, w_px):
        self.k = SLIDE_W / w_px  # EMU / px
        self.pt = 960.0 / w_px   # pt / px (13.333in × 72pt)

    def __call__(self, v):
        return Emu(int(round(v * self.k)))


def rgb(h):
    return RGBColor.from_string(h.lstrip("#").upper())


def face(run):
    w = int(run.get("weight") or (700 if run.get("bold") else 400))
    w = min(WEIGHT_FACE, key=lambda k: abs(k - w))
    return WEIGHT_FACE[w]


def set_face(rPr, name):
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", name)


def name_shape(shape, it, n):
    label = SHAPE_NAMES.get(it.get("cls") or "", it.get("cls") or "")
    if label:
        shape.name = f"{label} {n}"


def strip_style(shape):
    """add_shape가 붙이는 기본 스타일(그림자, 테마 색 테두리)을 뗀다"""
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)


def set_bullet(para, bu, sc, first_size, marL, hang):
    pPr = para._p.get_or_add_pPr()
    pPr.set("marL", str(int(sc(marL))))
    pPr.set("indent", str(-int(sc(hang))))
    for tag in ("a:buClr", "a:buSzPct", "a:buFont", "a:buChar", "a:buAutoNum", "a:buNone"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    if bu.get("color"):
        sf = etree.SubElement(pPr, qn("a:buClr"))
        etree.SubElement(sf, qn("a:srgbClr")).set("val", bu["color"].lstrip("#").upper())
    if bu.get("size") and first_size:
        etree.SubElement(pPr, qn("a:buSzPct")).set("val", str(int(min(400, max(25, bu["size"] / first_size * 100)) * 1000)))
    if bu.get("auto"):
        an = etree.SubElement(pPr, qn("a:buAutoNum"))
        an.set("type", bu["auto"])
        if bu.get("start", 1) > 1:
            an.set("startAt", str(bu["start"]))
    else:
        etree.SubElement(pPr, qn("a:buFont")).set("typeface", FAMILY)
        etree.SubElement(pPr, qn("a:buChar")).set("char", bu["char"])


def fill_runs(tf, paras, sc, first=True):
    for pi, p in enumerate(paras):
        para = tf.paragraphs[0] if (first and pi == 0) else tf.add_paragraph()
        para.alignment = ALIGN.get(p.get("align", "left"), PP_ALIGN.LEFT)
        # 한글을 어절 단위로 줄바꿈한다 (PowerPoint의 "한글 단어 잘림 허용"을 끈 상태). HTML의 word-break:keep-all과 같다.
        # 아래 fill_runs에서 글자 언어를 ko-KR로 지정하는 것과 함께 있어야 동작한다
        para._p.get_or_add_pPr().set("latinLnBrk", "0")
        lh = p.get("lineHeight")
        if lh:
            para.line_spacing = Pt(round(lh * sc.pt * 4) / 4)
        para.space_before = Pt(round(p.get("sb", 0) * sc.pt * 2) / 2)
        para.space_after = Pt(round(p.get("sa", 0) * sc.pt * 2) / 2)
        bu = p.get("bullet")
        # HTML에서 "• "를 글자로 친 줄도 PowerPoint에서는 글머리표 서식으로 바꾼다
        if not bu and p["runs"]:
            m = re.match(r"^([•·–])\s+", p["runs"][0]["text"])
            if m:
                p = dict(p, runs=[dict(p["runs"][0], text=p["runs"][0]["text"][m.end():])] + p["runs"][1:])
                bu = {"char": m.group(1), "color": p["runs"][0].get("color"), "size": p["runs"][0]["size"]}
        first_size = p["runs"][0]["size"] if p["runs"] else 12
        if bu:
            hang = p.get("hang") or first_size * 1.0
            set_bullet(para, bu, sc, first_size, p.get("ml", hang), hang)
        elif p.get("ml"):
            para._p.get_or_add_pPr().set("marL", str(int(sc(p["ml"]))))
        for r in p["runs"]:
            run = para.add_run()
            run.text = r["text"]
            f = run.font
            f.size = Pt(round(r["size"] * sc.pt * 4) / 4)
            name, bold = face(r)
            f.bold, f.italic = bold, r.get("italic", False)
            if r.get("underline"):
                f.underline = True
            f.color.rgb = rgb(r.get("color", "#000000"))
            f.name = name
            rPr = run._r.get_or_add_rPr()
            set_face(rPr, name)
            # 글자 언어가 한국어여야 PowerPoint가 어절 단위로 줄을 바꾼다 (실험으로 확인: 언어가 없으면 latinLnBrk와 관계없이 낱말 중간에서 끊긴다)
            rPr.set("lang", "ko-KR")
            rPr.set("altLang", "en-US")


def add_rect(slide, it, sc, n):
    x, y, w, h = sc(it["x"]), sc(it["y"]), sc(it["w"]), sc(it["h"])
    if w <= 0 or h <= 0:
        return
    if it.get("poly"):
        pts = [(sc(it["x"] + px), sc(it["y"] + py)) for px, py in it["poly"]]
        fb = slide.shapes.build_freeform(pts[0][0], pts[0][1], scale=1.0)
        fb.add_line_segments(pts[1:], close=True)
        shp = fb.convert_to_shape()
    elif it.get("ellipse") or (it.get("radius", 0) >= min(it["w"], it["h"]) / 2 - 0.5 and abs(it["w"] - it["h"]) <= 1):
        shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y, w, h)
    elif it.get("radius", 0) > 0.5:
        shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
        shp.adjustments[0] = min(0.5, it["radius"] / max(1, min(it["w"], it["h"])))
    else:
        shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    strip_style(shp)
    name_shape(shp, it, n)
    if it.get("fill"):
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(it["fill"])
        a = it.get("alpha", 1)
        if a < 0.99:
            clr = shp.fill._xPr.find(".//" + qn("a:srgbClr"))
            etree.SubElement(clr, qn("a:alpha")).set("val", str(int(a * 100000)))
    else:
        shp.fill.background()
    ln = it.get("line")
    if ln:
        shp.line.color.rgb = rgb(ln["color"])
        shp.line.width = Pt(max(0.25, ln["w"] * sc.pt))
        if ln.get("dash"):
            from pptx.enum.dml import MSO_LINE
            shp.line.dash_style = MSO_LINE.DASH if ln["dash"] == "dashed" else MSO_LINE.ROUND_DOT
    else:
        shp.line.fill.background()
    shp.text_frame.text = ""


def add_line(slide, it, sc):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, sc(it["x1"]), sc(it["y1"]), sc(it["x2"]), sc(it["y2"]))
    strip_style(c)
    c.line.color.rgb = rgb(it["color"])
    c.line.width = Pt(max(0.25, it["w"] * sc.pt))
    if it.get("dash"):
        from pptx.enum.dml import MSO_LINE
        c.line.dash_style = MSO_LINE.DASH if it["dash"] == "dashed" else MSO_LINE.ROUND_DOT


def text_box_geometry(it):
    """글상자의 (x, y, w, h, 줄바꿈 여부)를 정한다."""
    if it.get("list"):
        return it["x"], it["y"], it["w"], it["h"], True
    single = it.get("lines", 1) <= len(it["paras"])
    align = it["paras"][0].get("align", "left") if it["paras"] else "left"
    # 블록 요소의 한 줄 글자: 블록 너비 전체를 글상자로 삼고 줄바꿈을 켠다.
    # 사용자가 PowerPoint에서 글자를 늘려도 블록 안에서 줄바꿈되어 배치가 유지된다.
    if single and not it.get("vert") and it.get("block") and it.get("bw", 0) > it["w"] * 1.12:
        return it["bx"], it["y"], it["bw"], it["h"], True
    # 그 밖의 한 줄 글자는 줄바꿈을 끈다 (글자 폭의 미세한 차이로 끝 글자가 다음 줄로 떨어지는 것을 막는다)
    w = it["w"] * (1.04 if single else 1.015) + 2
    x = it["x"]
    if single and align == "center":
        x -= (w - it["w"]) / 2
    elif single and align in ("right", "end"):
        x -= (w - it["w"])
    return x, it["y"], w, it["h"], not single


def add_text(slide, it, sc, n):
    x, y, w, h, wrap = text_box_geometry(it)
    if it.get("role") == "title" and slide.shapes.title is not None:
        tb = slide.shapes.title
        # 헤드메시지는 제목 영역 전체(2줄 높이)를 개체 틀로 삼는다
        tb.left, tb.top, tb.width, tb.height = sc(it.get("bx", x)), sc(it.get("elTop", y)), sc(it.get("bw", w)), sc(max(it.get("zone", h), h))
        wrap = True
    else:
        tb = slide.shapes.add_textbox(sc(x), sc(y), sc(w), sc(max(h, 1)))
        name_shape(tb, it, n)
    tf = tb.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = wrap
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = MSO_ANCHOR.TOP
    if it.get("vert"):
        tf._txBody.find(qn("a:bodyPr")).set("vert", "eaVert")
    fill_runs(tf, it["paras"], sc)


def cell_border(cell, side, b, sc):
    tcPr = cell._tc.get_or_add_tcPr()
    tag = {"t": "a:lnT", "r": "a:lnR", "b": "a:lnB", "l": "a:lnL"}[side]
    old = tcPr.find(qn(tag))
    if old is not None:
        tcPr.remove(old)
    ln = etree.Element(qn(tag))
    if b:
        ln.set("w", str(int(max(0.25, b["w"] * sc.pt) * 12700)))
        sf = etree.SubElement(ln, qn("a:solidFill"))
        etree.SubElement(sf, qn("a:srgbClr")).set("val", b["color"].lstrip("#").upper())
        if b.get("dash"):
            etree.SubElement(ln, qn("a:prstDash")).set("val", "dash" if b["dash"] == "dashed" else "sysDot")
    else:
        ln.set("w", "0")
        etree.SubElement(ln, qn("a:noFill"))
    # a:lnL, lnR, lnT, lnB는 채움보다 앞에, 이 순서로 와야 한다 (스키마 순서)
    order = ["a:lnL", "a:lnR", "a:lnT", "a:lnB"]
    idx = 0
    for child in list(tcPr):
        if child.tag in [qn(o) for o in order[:order.index(tag)]]:
            idx = list(tcPr).index(child) + 1
    tcPr.insert(idx, ln)


def add_table(slide, it, sc, n):
    nr, nc = len(it["rows"]), len(it["cols"])
    if not nr or not nc:
        return
    gf = slide.shapes.add_table(nr, nc, sc(it["x"]), sc(it["y"]), sc(it["w"]), sc(it["h"]))
    gf.name = f"표 {n}"
    tbl = gf.table
    tblPr = tbl._tbl.tblPr
    for attr in ("firstRow", "bandRow", "firstCol", "lastRow", "lastCol", "bandCol"):
        tblPr.set(attr, "0")
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is None:
        sid = etree.SubElement(tblPr, qn("a:tableStyleId"))
    sid.text = NO_STYLE_TABLE
    for i, w in enumerate(it["cols"]):
        tbl.columns[i].width = sc(w)
    for i, h in enumerate(it["rows"]):
        tbl.rows[i].height = sc(h)
    for row in tbl.rows:  # 병합으로 가려지는 셀까지 기본 18pt를 없앤다
        for cell in row.cells:
            cell.text_frame.paragraphs[0]._p.get_or_add_endParaRPr().set("sz", "100")
    for c in it["cells"]:
        r0, c0, rs, cs = c["r"], c["c"], c["rs"], c["cs"]
        if r0 >= nr or c0 >= nc:
            continue
        cell = tbl.cell(r0, c0)
        if rs > 1 or cs > 1:
            cell.merge(tbl.cell(min(nr, r0 + rs) - 1, min(nc, c0 + cs) - 1))
        pt, pr, pb, pl = c["pad"]
        cell.margin_top, cell.margin_right, cell.margin_bottom, cell.margin_left = sc(pt), sc(pr), sc(pb), sc(pl)
        cell.vertical_anchor = VALIGN.get(c.get("valign"), MSO_ANCHOR.MIDDLE)
        if c.get("fill"):
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(c["fill"])
        else:
            cell.fill.background()
        for side, b in zip("trbl", c["borders"]):
            cell_border(cell, side, b, sc)
        # 병합된 칸은 첫 칸에만 서식이 들어가므로, 이어진 칸에도 같은 채움과 괘선을 넣는다 (간트 막대가 두 가지 농도로 보이는 것을 막는다)
        if rs > 1 or cs > 1:
            for rr in range(r0, min(nr, r0 + rs)):
                for cc in range(c0, min(nc, c0 + cs)):
                    if (rr, cc) == (r0, c0):
                        continue
                    sp = tbl.cell(rr, cc)
                    if c.get("fill"):
                        sp.fill.solid()
                        sp.fill.fore_color.rgb = rgb(c["fill"])
                    else:
                        sp.fill.background()
                    bt, br_, bb, bl = c["borders"]
                    cell_border(sp, "t", bt if rr == r0 else None, sc)
                    cell_border(sp, "b", bb if rr == r0 + rs - 1 else None, sc)
                    cell_border(sp, "l", None, sc)
                    cell_border(sp, "r", br_ if cc == c0 + cs - 1 else None, sc)
        if c.get("vert"):
            cell._tc.get_or_add_tcPr().set("vert", "eaVert")
        tf = cell.text_frame
        tf.word_wrap = True
        if c["paras"]:
            fill_runs(tf, c["paras"], sc)
        else:
            # 빈 셀(도형을 겹쳐 놓는 셀 포함)은 기본 18pt 때문에 행이 늘어나지 않도록 작은 글자 크기를 준다
            h_pt = it["rows"][r0] * sc.pt - (pt + pb) * sc.pt
            size = max(1, min(8, h_pt * 0.7))
            endp = tf.paragraphs[0]._p.get_or_add_endParaRPr()
            endp.set("sz", str(int(size * 100)))
            set_face(endp, FAMILY)


def patch_theme(prs, tokens):
    """테마 글꼴을 Pretendard로, 테마 색을 덱의 토큰으로 바꾼다."""
    part = prs.slide_masters[0].part.part_related_by(RT.THEME)
    root = etree.fromstring(part.blob)
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    for grp in ("majorFont", "minorFont"):
        for el in root.xpath(f".//a:{grp}", namespaces=ns):
            for tag in ("latin", "ea", "cs"):
                t = el.find(qn(f"a:{tag}"))
                if t is not None:
                    t.set("typeface", "Pretendard SemiBold" if grp == "majorFont" else FAMILY)
            for f in el.findall(qn("a:font")):
                if f.get("script") in ("Hang", "Jpan", "Hans", "Hant"):
                    f.set("typeface", "Pretendard SemiBold" if grp == "majorFont" else FAMILY)
    slot = {"dk2": "primary", "accent1": "accent", "accent2": "primary", "accent3": "bar", "accent4": "muted",
            "accent5": "brand", "accent6": "neg", "lt2": "soft"}
    for name, tok in slot.items():
        hexv = tokens.get(tok)
        el = root.find(f".//a:clrScheme/a:{name}", namespaces=ns)
        if hexv is None or el is None:
            continue
        for ch in list(el):
            el.remove(ch)
        etree.SubElement(el, qn("a:srgbClr")).set("val", hexv.lstrip("#").upper())
    scheme = root.find(".//a:clrScheme", namespaces=ns)
    if scheme is not None:
        scheme.set("name", "bmc-ppt")
    fonts = root.find(".//a:fontScheme", namespaces=ns)
    if fonts is not None:
        fonts.set("name", "Pretendard")
    part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def dump(html, td):
    here = Path(__file__).resolve().parent
    r = subprocess.run(["node", str(here / "html_dump.mjs"), str(Path(html).resolve()), td],
                       capture_output=True, text=True, timeout=600)
    if r.returncode != 0:
        sys.exit(f"html_dump.mjs 실패: {r.stderr.strip()[:400]}")
    return json.loads((Path(td) / "dump.json").read_text())


def convert(html, out):
    with tempfile.TemporaryDirectory() as td:
        data = dump(html, td)
        sc = Scale(data["slideW"])
        prs = Presentation()
        prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
        patch_theme(prs, data.get("tokens") or {})
        prs.core_properties.title = data.get("title") or ""
        prs.core_properties.author = ""
        prs.core_properties.last_modified_by = ""
        title_only, blank = prs.slide_layouts[5], prs.slide_layouts[6]
        images = 0
        for s in data["slides"]:
            items = s["items"]
            has_title = any(it["t"] == "text" and it.get("role") == "title" for it in items)
            slide = prs.slides.add_slide(title_only if has_title else blank)
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = rgb(s["bg"])
            # 겹침 순서: 채움·선·그림·표는 HTML 순서대로, 글자는 맨 위
            n = 0
            for it in items:
                n += 1
                t = it["t"]
                if t == "rect":
                    add_rect(slide, it, sc, n)
                elif t == "line":
                    add_line(slide, it, sc)
                elif t == "img":
                    p = Path(td) / it["file"]
                    # img 태그가 가리키는 원본이 이 컴퓨터의 PNG·JPG면 원본을 넣는다 (화면에서 찍은 것보다 해상도가 높다)
                    src = it.get("src") or ""
                    if src.startswith("file://"):
                        orig = Path(unquote(urlparse(src).path))
                        if orig.exists() and orig.suffix.lower() in (".png", ".jpg", ".jpeg", ".gif"):
                            p = orig
                    if p.exists():
                        pic = slide.shapes.add_picture(str(p), sc(it["x"]), sc(it["y"]), sc(it["w"]), sc(it["h"]))
                        pic.name = f"그림 {n}" + (f" {it['alt']}" if it.get("alt") else "")
                        images += 1
                elif t == "table":
                    add_table(slide, it, sc, n)
                elif t == "chart":
                    try:
                        add_chart(slide, it, sc, n)
                    except ValueError as e:
                        sys.exit(f"{s['n']}장의 기본 차트를 만들지 못했다: {e}")
            for it in items:
                n += 1
                if it["t"] == "text" and it["paras"]:
                    add_text(slide, it, sc, n)
        prs.save(out)
    return len(data["slides"]), images


def main():
    args = sys.argv[1:]
    if not args:
        sys.exit(__doc__)
    html = args[0]
    out = args[1] if len(args) > 1 else str(Path(html).with_suffix(".pptx"))
    n, images = convert(html, out)
    print(f"{n}장 → {out}")
    if images:
        print(f"참고: 그림으로 변환된 요소가 {images}개 있다 (SVG·img·배경 이미지). 그림 안의 글자와 수치는 편집할 수 없다.")


if __name__ == "__main__":
    main()
