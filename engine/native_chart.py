"""html_dump.mjs가 기록한 chart 항목(.nchart)을 PowerPoint 기본 차트로 만든다. html_to_pptx.py가 부른다.

- 수치는 차트 안의 엑셀 데이터로 들어간다. PowerPoint에서 차트를 선택하고 '데이터 편집'으로 고친다.
- 플롯 영역과 범례의 위치, 색은 design/nchart.js가 미리보기를 그리며 data-layout에 적은 값을 그대로 쓴다.
  그래서 HTML 미리보기와 PowerPoint 화면이 같은 자리에 같은 색으로 그려진다.
- 값 라벨은 c:dLbls를 직접 만들어 넣는다. 조각마다 글자색이 달라야 해서(짙은 조각은 흰 글자) python-pptx의 API로는 부족하다.
"""
from lxml import etree
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_MARKER_STYLE, XL_TICK_MARK
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Pt

FACE = {400: "Pretendard", 500: "Pretendard Medium", 600: "Pretendard SemiBold"}


def number_format(spec):
    d = min(6, max(0, int(spec.get("decimals") or 0)))
    f = "#,##0" + ("." + "0" * d if d else "")
    unit = spec.get("unit") or ""
    return f + (f'"{unit}"' if unit else "")


def _hex(h):
    return h.lstrip("#").upper()


def _style_font(font, px, color, weight, sc):
    """python-pptx Font에 크기, 색, Pretendard 굵기별 글꼴 이름을 준다 (한글 글꼴 자리 ea까지)."""
    font.size = Pt(round(px * sc.pt * 4) / 4)
    font.color.rgb = RGBColor.from_string(_hex(color))
    font.bold = False
    rPr = font._rPr
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", FACE[weight])
    rPr.set("lang", "ko-KR")


def _txpr(px, color, weight, sc):
    sz = int(round(px * sc.pt * 4) / 4 * 100)
    f = FACE[weight]
    return (f'<c:txPr><a:bodyPr/><a:lstStyle/><a:p><a:pPr><a:defRPr sz="{sz}" b="0">'
            f'<a:solidFill><a:srgbClr val="{_hex(color)}"/></a:solidFill>'
            f'<a:latin typeface="{f}"/><a:ea typeface="{f}"/><a:cs typeface="{f}"/></a:defRPr></a:pPr>'
            f'<a:endParaRPr lang="ko-KR"/></a:p></c:txPr>')


def _flags(val, cat, pos=None, sep=None):
    s = f'<c:dLblPos val="{pos}"/>' if pos else ""
    s += (f'<c:showLegendKey val="0"/><c:showVal val="{int(val)}"/><c:showCatName val="{int(cat)}"/>'
          '<c:showSerName val="0"/><c:showPercent val="0"/><c:showBubbleSize val="0"/>')
    if sep is not None:
        s += f"<c:separator>{sep}</c:separator>"
    return s


def _dlbls(points, group, fmt):
    """points: [(idx, txPr, flags)], group: (txPr, flags) 계열 전체의 기본값"""
    body = "".join(f'<c:dLbl><c:idx val="{i}"/><c:delete val="1"/></c:dLbl>' if tx is None else
                   f'<c:dLbl><c:idx val="{i}"/><c:numFmt formatCode="{_esc(fmt)}" sourceLinked="0"/>'
                   f'<c:spPr><a:noFill/><a:ln><a:noFill/></a:ln></c:spPr>{tx}{fl}</c:dLbl>' for i, tx, fl in points)
    tx, fl = group
    body += (f'<c:numFmt formatCode="{_esc(fmt)}" sourceLinked="0"/><c:spPr><a:noFill/><a:ln><a:noFill/></a:ln></c:spPr>'
             f"{tx}{fl}")
    return etree.fromstring(f"<c:dLbls {nsdecls('c', 'a')}>{body}</c:dLbls>")


def _esc(s):
    return s.replace("&", "&amp;").replace('"', "&quot;").replace("<", "&lt;")


def _put_dlbls(ser, dl):
    old = ser.find(qn("c:dLbls"))
    if old is not None:
        ser.remove(old)
    anchor = ser.find(qn("c:cat"))
    anchor.addprevious(dl)


def _manual(parent, box, W, H, inner, after=None):
    """c:layout에 manualLayout(차트 전체 대비 비율)을 넣는다. parent는 c:plotArea 또는 c:legend"""
    lay = parent.find(qn("c:layout"))
    if lay is None:
        lay = etree.SubElement(parent, qn("c:layout"))
        if after is not None:
            parent.find(qn(after)).addnext(lay)
        else:
            parent.insert(0, lay)
    for ch in list(lay):
        lay.remove(ch)
    if not (box["w"] > 0 and box["h"] > 0 and -0.5 <= box["x"] and -0.5 <= box["y"]
            and box["x"] + box["w"] <= W + 0.5 and box["y"] + box["h"] <= H + 0.5):
        raise ValueError(f"차트 배치가 차트 영역을 벗어난다: {box} (영역 {W}×{H})")
    tgt = '<c:layoutTarget val="inner"/>' if inner else ""
    ml = (f"<c:manualLayout {nsdecls('c')}>{tgt}<c:xMode val=\"edge\"/><c:yMode val=\"edge\"/>"
          f'<c:x val="{box["x"] / W:.5f}"/><c:y val="{box["y"] / H:.5f}"/><c:w val="{box["w"] / W:.5f}"/><c:h val="{box["h"] / H:.5f}"/>'
          "</c:manualLayout>")
    lay.append(etree.fromstring(ml))


def _clear_box(parent, before):
    """parent(c:chartSpace 또는 c:plotArea)의 spPr을 채움 없음, 테두리 없음으로 바꾼다. 스키마 순서에 맞는 자리에 넣는다"""
    old = parent.find(qn("c:spPr"))
    if old is not None:
        parent.remove(old)
    sp = etree.fromstring(f"<c:spPr {nsdecls('c', 'a')}><a:noFill/><a:ln><a:noFill/></a:ln></c:spPr>")
    nxt = next((parent.find(qn(t)) for t in before if parent.find(qn(t)) is not None), None)
    if nxt is not None:
        nxt.addprevious(sp)
    else:
        parent.append(sp)


def _no_line(fmt):
    fmt.line.fill.background()


def add_chart(slide, it, sc, n):
    spec, L = it["spec"], it["layout"]
    T = L["tokens"]
    typ = spec["type"]
    W, H = L["W"], L["H"]
    fmt = number_format(spec)
    # python-pptx는 이 서식을 이스케이프하지 않고 XML에 넣는다(단위에 &가 있으면 깨진다). 값 라벨과 축에 서식을 따로 주므로 여기서는 비운다
    cd = CategoryChartData()
    cd.categories = [str(c) for c in spec["categories"]]
    for s in spec["series"]:
        cd.add_series(str(s["name"]), [float(v) for v in s["values"]])
    xl = {"doughnut": XL_CHART_TYPE.DOUGHNUT, "pie": XL_CHART_TYPE.PIE, "line": XL_CHART_TYPE.LINE_MARKERS,
          "area": XL_CHART_TYPE.AREA_STACKED if L.get("stacked") else XL_CHART_TYPE.AREA}[typ]
    gf = slide.shapes.add_chart(xl, sc(it["x"]), sc(it["y"]), sc(it["w"]), sc(it["h"]), cd)
    gf.name = f"차트 {n}"
    ch = gf.chart
    ch.has_title = False
    _style_font(ch.font, 12, T["ink"], 400, sc)
    cs = ch._chartSpace
    # 차트 영역과 플롯 영역의 테두리와 채움을 없앤다 (슬라이드 바탕이 그대로 보이게)
    _clear_box(cs, before=("c:txPr", "c:externalData", "c:printSettings", "c:userShapes", "c:extLst"))
    _clear_box(cs.chart.plotArea, before=("c:extLst",))
    _manual(cs.chart.plotArea, L["plot"], W, H, inner=True)
    plot = ch.plots[0]
    cols = L["colors"]
    hi = spec.get("highlight")

    if typ in ("doughnut", "pie"):
        plot.vary_by_categories = True
        ser = plot.series[0]
        for i, pt in enumerate(ser.points):
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = RGBColor.from_string(_hex(cols[i]))
            pt.format.line.color.rgb = RGBColor.from_string(_hex(T["slide-bg"]))
            pt.format.line.width = Pt(1.5 * sc.pt)
        if typ == "doughnut":
            hs = cs.find(".//" + qn("c:holeSize"))
            if hs is not None:
                hs.set("val", "60")
            vals = spec["series"][0]["values"]
            pts = [(i, None, None) if vals[i] <= 0 else (i, _txpr(13, L["labelColors"][i], 600, sc), _flags(1, 0)) for i in range(len(cols))]
            _put_dlbls(ser._element, _dlbls(pts, (_txpr(13, T["primary"], 600, sc), _flags(1, 0)), fmt))
        else:
            vals = spec["series"][0]["values"]
            pts = [(i, None, None) if vals[i] <= 0 else (i, _txpr(13, T["accent"] if i == hi else T["primary"], 600, sc), _flags(1, 1, "outEnd", "\n"))
                   for i in range(len(cols))]
            _put_dlbls(ser._element, _dlbls(pts, (_txpr(13, T["primary"], 600, sc), _flags(1, 1, "outEnd", "\n")), fmt))
        ch.has_legend = False
        if L.get("legend"):
            ch.has_legend = True
            ch.legend.position = XL_LEGEND_POSITION.RIGHT
            ch.legend.include_in_layout = False
            _style_font(ch.legend.font, 13, T["ink"], 400, sc)
            _manual(ch.legend._element, L["legend"], W, H, inner=False, after="c:legendPos")
    else:
        line = typ == "line"
        mode = L.get("labels", "none")
        for i, s in enumerate(plot.series):
            c = RGBColor.from_string(_hex(cols[i]))
            label_c = T["accent"] if i == hi else T["primary"]
            if line:
                s.smooth = False
                s.format.line.color.rgb = c
                s.format.line.width = Pt(3 * sc.pt)
                s.marker.style = XL_MARKER_STYLE.CIRCLE
                s.marker.size = max(2, round(9 * sc.pt))
                s.marker.format.fill.solid()
                s.marker.format.fill.fore_color.rgb = c
                s.marker.format.line.color.rgb = c
                n_pts = len(spec["categories"])
                if mode == "all":
                    _put_dlbls(s._element, _dlbls([], (_txpr(12.5, label_c, 600, sc), _flags(1, 0, "t")), fmt))
                elif mode == "last":
                    pts = [(n_pts - 1, _txpr(12.5, label_c, 600, sc), _flags(1, 0, "t"))]
                    _put_dlbls(s._element, _dlbls(pts, (_txpr(12.5, label_c, 600, sc), _flags(0, 0)), fmt))
            else:
                s.format.fill.solid()
                s.format.fill.fore_color.rgb = c
                _no_line(s.format)
        va, ca = ch.value_axis, ch.category_axis
        ymax, major = float(L["ymax"]), float(L["major"])
        if not (0 < major <= 1e10 and ymax > 0):
            raise ValueError(f"세로축 눈금 {major}이 PowerPoint 허용 범위(0 초과 100억 이하) 밖이다. 단위를 키워(예: 원 → 억 원) 값을 줄인다")
        va.minimum_scale, va.maximum_scale, va.major_unit = 0, ymax, major
        va.has_major_gridlines = True
        va.major_gridlines.format.line.color.rgb = RGBColor.from_string(_hex(T["hair"]))
        va.major_gridlines.format.line.width = Pt(1 * sc.pt)
        va.major_tick_mark = va.minor_tick_mark = XL_TICK_MARK.NONE
        _no_line(va.format)
        va.tick_labels.number_format = fmt
        va.tick_labels.number_format_is_linked = False
        _style_font(va.tick_labels.font, 11.5, T["muted"], 400, sc)
        ca.major_tick_mark = ca.minor_tick_mark = XL_TICK_MARK.NONE
        ca.format.line.color.rgb = RGBColor.from_string(_hex(T["primary"]))
        ca.format.line.width = Pt(1.5 * sc.pt)
        _style_font(ca.tick_labels.font, 12, T["muted"], 400, sc)
        cb = cs.find(".//" + qn("c:valAx") + "/" + qn("c:crossBetween"))
        if cb is not None:
            cb.set("val", "between" if L.get("between") else "midCat")
        ch.has_legend = False
        if L.get("legend"):
            ch.has_legend = True
            ch.legend.position = XL_LEGEND_POSITION.TOP
            ch.legend.include_in_layout = False
            _style_font(ch.legend.font, 12, T["ink"], 400, sc)
            _manual(ch.legend._element, L["legend"], W, H, inner=False, after="c:legendPos")
    return gf
